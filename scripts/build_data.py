import openpyxl, csv, collections, json, math
import sys
SRC=(sys.argv[1] if len(sys.argv)>1 else "raw")+"/"   # folder holding the original .xlsx files (run from repo root)
OUT="data/"
YEARS=[2009,2012,2015,2018,2021,2024]
STATE_ABBR={'New South Wales':'NSW','Victoria':'VIC','Queensland':'QLD','South Australia':'SA','Western Australia':'WA','Tasmania':'TAS','Northern Territory':'NT','Australian Capital Territory':'ACT'}
RENAME={25250:24700}   # Moreland -> Merri-bek (code changed in ASGS 2023 LGA edition)

def num(v):
    if isinstance(v,(int,float)) and not isinstance(v,bool): return v
    if isinstance(v,str):
        s=v.strip()
        if s.startswith('≥'): return float(s[1:])        # bounded value: use the bound
        if s.startswith('≤'): return float(s[1:])
    return None

def rows(fn, sheet):
    wb=openpyxl.load_workbook(SRC+fn, data_only=True, read_only=True)
    ws=wb[sheet]; ws.reset_dimensions()
    return [list(r)+[None]*50 for r in ws.iter_rows(values_only=True)]

# ---------- 1. AEDC by council, long format (charts 1-4) ----------
cent={}
with open("scripts/centroids.csv") as f:
    for r in csv.DictReader(f): cent[int(r['LGA_CODE24'])]=(float(r['lon']),float(r['lat']))
out=[]; state=None
for r in rows("lga-2009-24.xlsx","LGA One or more")[6:]:
    b,c=r[1],r[2]
    if isinstance(b,str) and b in STATE_ABBR: state=STATE_ABBR[b]; continue
    if isinstance(b,(int,float)) and float(b).is_integer():
        code=RENAME.get(int(b),int(b))
        lon,lat=cent.get(code,(None,None))
        if lon is None: continue                    # Norfolk Island etc. (not on the map)
        for k,y in enumerate(YEARS):
            valid=num(r[3+k]); n=num(r[9+2*k]); pct=num(r[10+2*k])
            out.append(dict(lga_code=code,lga_name=c,state=state,year=y,
                            children_assessed=int(valid) if valid is not None else '',
                            vulnerable_n=int(n) if n is not None else '',
                            vulnerable_pct=round(pct,1) if pct is not None else '',
                            lon=lon,lat=lat))
# Hexagon bins (chart 3): pointy-top hexes, 100 km flat-to-flat, built in the same Albers projection as the maps.
# The grid origin is shifted 40 km east / 20 km north so Sydney, Melbourne and Brisbane each split across 3+ hexes.
from pyproj import Transformer
AEA="+proj=aea +lat_1=-18 +lat_2=-36 +lat_0=0 +lon_0=133.5 +datum=WGS84 +units=m"
fwd=Transformer.from_crs("EPSG:4326",AEA,always_xy=True); inv=Transformer.from_crs(AEA,"EPSG:4326",always_xy=True)
HEX_W=100_000; HEX_S=HEX_W/3**.5; HEX_OX,HEX_OY=40_000,20_000
def hex_of(lon,lat):
    x,y=fwd.transform(lon,lat); x-=HEX_OX; y-=HEX_OY
    q=(3**.5/3*x-y/3)/HEX_S; r=(2/3*y)/HEX_S; z=-q-r
    rq,rr,rz=round(q),round(r),round(z)                           # cube rounding
    dq,dr,dz=abs(rq-q),abs(rr-r),abs(rz-z)
    if dq>dr and dq>dz: rq=-rr-rz
    elif dr>dz: rr=-rq-rz
    return rq,rr
def hex_center(q,r): return HEX_S*3**.5*(q+r/2)+HEX_OX, HEX_S*1.5*r+HEX_OY
hexes={}
for o in out:
    q,r=hex_of(o['lon'],o['lat']); o['hex_id']=f"{q}_{r}"; hexes[o['hex_id']]=(q,r)
feats=[]
for hid,(q,r) in sorted(hexes.items()):
    cx,cy=hex_center(q,r)
    ring=[inv.transform(cx+HEX_S*math.cos(math.radians(30+60*i)),cy+HEX_S*math.sin(math.radians(30+60*i))) for i in range(6)]
    ring=[[round(a,5),round(b,5)] for a,b in ring][::-1]            # clockwise in lon/lat, the winding d3-geo needs
    ring.append(ring[0])
    clon,clat=inv.transform(cx,cy)
    feats.append(dict(type="Feature",properties=dict(hex_id=hid,cx=round(clon,4),cy=round(clat,4)),
                      geometry=dict(type="Polygon",coordinates=[ring])))
with open(OUT+"hex_grid.geojson","w") as f: json.dump(dict(type="FeatureCollection",features=feats),f,separators=(',',':'))
print("hex_grid.geojson", len(feats), "hexes")

with open(OUT+"aedc_lga.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(out[0])); w.writeheader(); w.writerows(out)
print("aedc_lga.csv", len(out), "rows,", len({o['lga_code'] for o in out}), "councils")

# ---------- 2. AEDC by council, 2024, vulnerable per domain (chart 1 dropdown) ----------
# One or more: n = r[19], % = r[20]. Domain sheets: vulnerable pairs start at r[33], so 2024 is r[43], r[44].
MEASURES=[('One or more domains','LGA One or more',19),('Physical health','LGA Health',43),
          ('Social competence','LGA Social',43),('Emotional maturity','LGA Emotional',43),
          ('Language & thinking','LGA Language',43),('Communication','LGA Communication',43)]
lgadom=[]
for measure,sheet,ci in MEASURES:
    state=None
    for r in rows("lga-2009-24.xlsx",sheet)[6:]:
        b,c=r[1],r[2]
        if isinstance(b,str) and b in STATE_ABBR: state=STATE_ABBR[b]; continue
        if isinstance(b,(int,float)) and float(b).is_integer():
            code=RENAME.get(int(b),int(b))
            if code not in cent: continue
            valid=num(r[8]); n=num(r[ci]); pct=num(r[ci+1])
            lgadom.append(dict(lga_code=code,lga_name=c,state=state,measure=measure,
                               children_assessed=int(valid) if valid is not None else '',
                               vulnerable_n=int(n) if n is not None else '',
                               vulnerable_pct=round(pct,1) if pct is not None else ''))
with open(OUT+"aedc_lga_domains.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(lgadom[0])); w.writeheader(); w.writerows(lgadom)
chk={(d['lga_name'],d['measure']):d['vulnerable_pct'] for d in lgadom}
print("aedc_lga_domains.csv", len(lgadom), "rows; check Albury DV1/Health, Brisbane DV1/Language:",
      chk[('Albury','One or more domains')], chk[('Albury','Physical health')],
      chk[('Brisbane','One or more domains')], chk[('Brisbane','Language & thinking')])

# ---------- 3. State trends: summary indicators (chart 5) + national domains (chart 6) ----------
wb=openpyxl.load_workbook(SRC+"state-and-territory-trends-(2009-2024).xlsx", data_only=True, read_only=True)
SHORT={'Developmentally vulnerable on one or more domains':'Vulnerable on 1+ domains',
       'Developmentally vulnerable on two or more domains':'Vulnerable on 2+ domains',
       'Developmentally on track on five domains':'On track on all 5 domains'}
summ=[]; dom=collections.defaultdict(lambda:[0,0])
for sn in wb.sheetnames[1:]:
    st=sn.split('. ')[-1].split('_')[0].replace('Table 11.','').replace('Table 15.','').strip()
    ws=wb[sn]; ws.reset_dimensions(); label=None
    for r in ws.iter_rows(values_only=True):
        r=list(r)+[None]*20
        if r[0] not in (None,''): label=r[0]
        if 'summary' in sn.lower():
            if isinstance(r[1],int) and isinstance(r[3],(int,float)):
                summ.append(dict(state=st,year=r[1],indicator=SHORT[label],pct=round(r[3],1),n=r[2],children=r[4],
                                 sig_2009_2024=(r[6] or '') if r[1]==2024 else '',sig_2021_2024=(r[7] or '') if r[1]==2024 else ''))
        else:
            if isinstance(r[1],int) and isinstance(r[6],(int,float)) and isinstance(r[8],(int,float)):
                dom[(label,r[1])][0]+=r[6]; dom[(label,r[1])][1]+=r[8]
# Australia = sum of n / sum of children across the 8 states (no significance test published for it)
aus=collections.defaultdict(lambda:[0,0])
for s in summ: aus[(s['indicator'],s['year'])][0]+=s['n']; aus[(s['indicator'],s['year'])][1]+=s['children']
summ+=[dict(state='Australia',year=y,indicator=ind,pct=round(100*n/c,1),n=n,children=c,sig_2009_2024='',sig_2021_2024='')
       for (ind,y),(n,c) in aus.items()]
with open(OUT+"aedc_states.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(summ[0])); w.writeheader(); w.writerows(summ)
print("aedc_states.csv", len(summ), "rows; states:", sorted({s['state'] for s in summ}))
print("  Australia on track 5, by year:", sorted((y,s['pct']) for s in summ for y in [s['year']] if s['state']=='Australia' and s['indicator']=='On track on all 5 domains'))
print("  Australia vulnerable 1+ n, by year:", sorted((s['year'],s['n']) for s in summ if s['state']=='Australia' and s['indicator']=='Vulnerable on 1+ domains'))
print("  on track 5 significance 2024 (2009 vs 2024 / 2021 vs 2024):", [(s['state'],s['sig_2009_2024'],s['sig_2021_2024']) for s in summ if s['year']==2024 and s['indicator']=='On track on all 5 domains' and s['state']!='Australia'])
DSHORT={'Physical health and wellbeing':'Physical health','Social competence':'Social competence',
        'Emotional maturity':'Emotional maturity','Language and cognitive skills (school-based)':'Language & thinking',
        'Communication skills and general knowledge':'Communication'}
domrows=[dict(domain=DSHORT[l],domain_full=l,year=y,vulnerable_pct=round(100*n/t,1)) for (l,y),(n,t) in sorted(dom.items())]
with open(OUT+"aedc_domains_national.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(domrows[0])); w.writeheader(); w.writerows(domrows)
print("domains 2021/2024:", [(d['domain'],d['year'],d['vulnerable_pct']) for d in domrows if d['year'] in (2021,2024)])

# National on track / at risk / vulnerable per domain, 2009 and 2024 (chart 0), summed over the 8 "all domains" sheets.
# r[0] = domain (first row of each block only), r[1] = year, r[2] on track n, r[4] at risk n, r[6] vulnerable n, r[8] total.
cat=collections.defaultdict(lambda:[0,0,0,0])
for sn in wb.sheetnames[1:]:
    if 'summary' in sn.lower(): continue
    ws=wb[sn]; ws.reset_dimensions(); label=None
    for r in ws.iter_rows(values_only=True):
        r=list(r)+[None]*20
        if r[0] not in (None,''): label=r[0]
        if r[1] in (2009,2024) and all(isinstance(r[i],(int,float)) for i in (2,4,6,8)):
            for j,i in enumerate((6,4,2,8)): cat[(label,r[1])][j]+=r[i]
DORDER=['Physical health','Social competence','Emotional maturity','Language & thinking','Communication']
CATS=['Vulnerable','At risk','On track']
catrows=[]
for (l,y),(v,ar,ot,t) in sorted(cat.items(),key=lambda kv:(DORDER.index(DSHORT[kv[0][0]]),kv[0][1])):
    for co,(c,n) in enumerate(zip(CATS,(v,ar,ot))):
        catrows.append(dict(domain=DSHORT[l],domain_order=DORDER.index(DSHORT[l]),year=y,category=c,category_order=co,n=n,total=t,pct=round(100*n/t,1)))
with open(OUT+"aedc_categories_national.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(catrows[0])); w.writeheader(); w.writerows(catrows)
natv={(d['domain'],d['year']):d['vulnerable_pct'] for d in domrows}
assert all(c['pct']==natv[(c['domain'],c['year'])] for c in catrows if c['category']=='Vulnerable'), "vulnerable % differs from aedc_domains_national.csv"
print("aedc_categories_national.csv 2024 (vulnerable / at risk / on track):")
for d in DORDER: print("  ",d,[c['pct'] for c in catrows if c['domain']==d and c['year']==2024])

# ---------- 4. SEIFA: most vs least disadvantaged, on track on five (chart 7) ----------
seifa=[]; nat=collections.defaultdict(lambda:[0,0]); state=None
for r in rows("seifa-2009-24.xlsx","SEIFA On track five"):
    b,c=r[1],r[2]
    if isinstance(b,str) and b in STATE_ABBR: state=STATE_ABBR[b]; continue
    if isinstance(c,str) and (c.startswith('Quintile 1') or c.startswith('Quintile 5')):
        q='Most disadvantaged' if c.startswith('Quintile 1') else 'Least disadvantaged'
        for k,y in enumerate(YEARS):
            valid=num(r[3+k]); n=num(r[9+2*k]); pct=num(r[10+2*k])
            if y in (2021,2024) and pct is not None:
                seifa.append(dict(area=state,group=q,year=y,on_track_pct=round(pct,1)))
                nat[(q,y)][0]+=n; nat[(q,y)][1]+=valid
for (q,y),(n,t) in nat.items():
    seifa.append(dict(area='Australia',group=q,year=y,on_track_pct=round(100*n/t,1)))
with open(OUT+"aedc_seifa.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(seifa[0])); w.writeheader(); w.writerows(seifa)
print("seifa national:", {k:round(100*v[0]/v[1],1) for k,v in nat.items()})

# ================= ABS NNPAS 2023 (charts 8–11) =================
def find(rs,label,start=0):
    for i,r in enumerate(rs[start:],start):
        if isinstance(r[0],str) and r[0].strip()==label: return i,r
    raise KeyError(label)
def findp(rs,prefix,start=0):
    for i,r in enumerate(rs[start:],start):
        if isinstance(r[0],str) and r[0].strip().startswith(prefix): return i,r
    raise KeyError(prefix)
def pnum(v): return v if isinstance(v,(int,float)) and not isinstance(v,bool) else None   # "np" -> None
SEXES=(('All','Children aged 5–17 years(c)'),('Boys','Males aged 5–17 years'),('Girls','Females aged 5–17 years'))
AGES=['5–8','9–11','12–14','15–17']
t5=rows("CPASSDC04–05.xlsx","Table 5.1_Means")
t12=rows("CPASSDC12.xlsx","Table 12.3_Proportions")
n2=rows("NNPASDC02.xlsx","Table 2.1_Means Persons"); m2=rows("NNPASDC02.xlsx","Table 2.2_MoEs Persons")

# ---------- 5. All 24-hour guidelines + each guideline, by age and sex (charts 8, 8b) ----------
g3=rows("CPASSDC01–02.xlsx","Table 2.3_Proportions"); g4=rows("CPASSDC01–02.xlsx","Table 2.4_MoEs")
assert [str(x) for x in g3[5][1:5]]==AGES
GUIDE=(('All guidelines','Met 24-Hour Movement Guidelines'),('Activity','Met physical activity recommendation'),
       ('Strength','Did muscle/bone strengthening activity on 3 or more days'),
       ('Screen time','Met sedentary screen time recommendation'),('Sleep','Met sleep recommendation'))
gl=[]
for sex,block in SEXES:
    b3=[i for i,r in enumerate(g3) if r[1]==block][0]; b4=[i for i,r in enumerate(g4) if r[1]==block][0]
    for go,(g,label) in enumerate(GUIDE):
        _,p=findp(g3,label,b3); _,m=findp(g4,label,b4)
        for k,a in enumerate(AGES):
            mo=pnum(m[1+k])
            gl.append(dict(sex=sex,age_group=a,age_order=k,guideline=g,guideline_order=go,pct=p[1+k],moe='' if mo is None else mo))
with open(OUT+"nnpas_guidelines.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(gl[0])); w.writeheader(); w.writerows(gl)
print("nnpas_guidelines.csv")
for sex in ('All','Boys','Girls'):
    for g,_ in GUIDE:
        if sex=='All' or g=='All guidelines':
            print("  ",sex,g,[x['pct'] for x in gl if x['sex']==sex and x['guideline']==g],"moe",[x['moe'] for x in gl if x['sex']==sex and x['guideline']==g])

# ---------- 6. Activity vs sedentary screen minutes per day, by age and sex (chart 9) ----------
t5r=rows("CPASSDC04–05.xlsx","Table 5.2_RSEs")                  # same layout as Table 5.1, values are RSE %
def act_screen_rows(tab,block):
    bi=[i for i,r in enumerate(tab) if r[1]==block][0]
    _,a=find(tab,'Total moderate or vigorous physical activity (incl. active transport)',bi)
    hi,_=findp(tab,'Average sedentary screen time per day',bi)
    _,s=find(tab,'All days',hi)                                  # the "All days" row inside the screen-time block
    return a,s
acts=[]
for sex,block in SEXES:
    a,s=act_screen_rows(t5,block); ar,sr=act_screen_rows(t5r,block)
    for k,ag in enumerate(AGES):
        acts.append(dict(sex=sex,age_group=ag,age_order=k,activity_min=round(a[1+k].total_seconds()/60),screen_min=round(s[1+k].total_seconds()/60),
                         activity_rse=pnum(ar[1+k]) if pnum(ar[1+k]) is not None else '',screen_rse=pnum(sr[1+k]) if pnum(sr[1+k]) is not None else ''))
with open(OUT+"nnpas_activity_screens.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(acts[0])); w.writeheader(); w.writerows(acts)
print("nnpas_activity_screens.csv")
for sex,_ in SEXES:
    print("  ",sex,"activity",[x['activity_min'] for x in acts if x['sex']==sex],"screens",[x['screen_min'] for x in acts if x['sex']==sex],
          "RSE activity",[x['activity_rse'] for x in acts if x['sex']==sex],"screens",[x['screen_rse'] for x in acts if x['sex']==sex])

# ---------- 6b. A child's day: sleep, activity, screens and the rest of 24 hours (day clock) ----------
mins=lambda v: round(v.total_seconds()/60)
t4=rows("CPASSDC04–05.xlsx","Table 4.1_Means")
assert str(t4[5][3]).startswith('Total 2–5')
_,sl4=find(t4,'Total sleep(d)(e)'); _,ac4=find(t4,'Total physical activity')
h4,_=findp(t4,'Average sedentary screen time per day'); _,sc4=find(t4,'All days',h4)
day_in=[('2–5',mins(sl4[3]),mins(ac4[3]),mins(sc4[3]),
         'Ages 2–5: activity includes light play and sleep includes naps (measured differently from older ages)')]
bi=[i for i,r in enumerate(t5) if r[1]==SEXES[0][1]][0]
_,sl5=find(t5,'Total sleep on night before interview(f)',bi)
_,ac5=find(t5,'Total moderate or vigorous physical activity (incl. active transport)',bi)
h5,_=findp(t5,'Average sedentary screen time per day',bi); _,sc5=find(t5,'All days',h5)
day_in+=[(a,mins(sl5[1+k]),mins(ac5[1+k]),mins(sc5[1+k]),'') for k,a in enumerate(AGES)]
day=[]
for k,(a,sl,ac,sc,note) in enumerate(day_in):
    for o,(act,m) in enumerate((('Sleep',sl),('Active',ac),('Screens',sc),('Everything else',1440-sl-ac-sc))):
        day.append(dict(age_group=a,age_order=k,activity=act,activity_order=o,minutes=m,label=f"{m//60}h {m%60:02d}m",note=note))
with open(OUT+"nnpas_day.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(day[0])); w.writeheader(); w.writerows(day)
print("nnpas_day.csv (sleep, active, screens, everything else):")
for a,*_ in day_in: print("  ",a,[x['minutes'] for x in day if x['age_group']==a])

# ---------- 7. Screen devices in the bedroom + sleep quality, by age (charts 10, 10b) ----------
AGES5=['2–5']+AGES
col12={str(h).replace('(c)','').replace('(b)',''):i for i,h in enumerate(t12[5]) if h}
DEV=(('Smartphone/watch','Smart phone or smart watch'),('Computer','Computer (including desktop or laptop)'),
     ('Television','Television (including DVDs, streaming services, free-to-air)'),('Tablet','Tablet'),
     ('Gaming console','Gaming console'))
t12m=rows("CPASSDC12.xlsx","Table 12.4_MoEs")                    # same rows and columns as Table 12.3
assert [str(x) for x in t12m[5][1:6]]==[str(x) for x in t12[5][1:6]]
moe12=lambda label:[(lambda v:'' if v is None else v)(pnum(findp(t12m,label)[1][col12[a]])) for a in AGES5]   # "np" -> blank
devs=[]
for d,label in DEV:
    _,r=findp(t12,label); devs.append((d,[r[col12[a]] for a in AGES5],moe12(label)))
devs.sort(key=lambda t:-t[1][-1])                                  # single devices ordered by the 15–17 value
TOT='Total with screen-based device located in bedroom'
_,r=findp(t12,TOT); devs.append(('Any screen device',[r[col12[a]] for a in AGES5],moe12(TOT)))
bed=[dict(age_group=a,age_order=k,device=d,device_order=o,pct=v[k],moe=m[k]) for o,(d,v,m) in enumerate(devs) for k,a in enumerate(AGES5)]
with open(OUT+"nnpas_bedroom_devices.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(bed[0])); w.writeheader(); w.writerows(bed)
print("nnpas_bedroom_devices.csv"); [print("  ",d,v,"moe",m) for d,v,m in devs]
RATE=(('Very good','Very good'),('Good','Good'),('Fair','Fair'),('Poor / very poor','Poor / Very poor'))
raw={rt:[findp(t12,label)[1][col12[a]] for a in AGES5] for rt,label in RATE}
rmoe={rt:moe12(label) for rt,label in RATE}
sq=[]
for k,a in enumerate(AGES5):
    tot=sum(raw[rt][k] for rt,_ in RATE)                           # "not known" sits in the published total; drop it
    for o,(rt,_) in enumerate(RATE):
        sq.append(dict(age_group=a,age_order=k,rating=rt,rating_order=o,pct_raw=raw[rt][k],pct=round(100*raw[rt][k]/tot,1),moe=rmoe[rt][k]))
with open(OUT+"nnpas_sleep_quality.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(sq[0])); w.writeheader(); w.writerows(sq)
print("nnpas_sleep_quality.csv")
for rt,_ in RATE: print("  ",rt,"raw",raw[rt],"normalised",[x['pct'] for x in sq if x['rating']==rt],"moe",rmoe[rt])

# ---------- 8. Saturated fat and free sugars vs dietary limits (chart 11) ----------
# Limits are fixed reference values: NHMRC Nutrient Reference Values (saturated + trans fat, no more than 10% of energy)
# and WHO Guideline: Sugars intake for adults and children, 2015 (free sugars below 10%, ideally below 5%).
NUTR=(('Saturated + trans fat','Saturated fat + trans fatty acids',10,''),('Free sugars','Free sugars',10,5))
hdr=n2[5]; nut=[]
for name,label,lim,ideal in NUTR:
    _,mr=find(n2,label); _,er=find(m2,label)
    for k,(a,lab,grp) in enumerate((('2–4','2–4','Children'),('5–11','5–11','Children'),('12–17','12–17','Children'),('18 years and over','Adults 18+','Adults'))):
        i=hdr.index(a); nut.append(dict(nutrient=name,age_group=lab,age_order=k,group=grp,pct_energy=mr[i],moe=er[i],limit=lim,ideal=ideal))
with open(OUT+"nnpas_nutrients.csv","w",newline="") as f:
    w=csv.DictWriter(f,fieldnames=list(nut[0])); w.writeheader(); w.writerows(nut)
print("nnpas_nutrients.csv")
for name,*_ in NUTR: print("  ",name,[(x['age_group'],x['pct_energy'],x['moe']) for x in nut if x['nutrient']==name])
